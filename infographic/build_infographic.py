"""Build the FerroFlux graduation infographic (SVG -> PNG/PDF/PPTX). Run: PYTHONPATH=<cairosvg libs> python3 build_infographic.py"""
import base64, io, os, re
from xml.sax.saxutils import escape
import qrcode
from PIL import Image, ImageChops, ImageFont, ImageOps
import cairosvg
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")
os.makedirs(OUT, exist_ok=True)

W, H = 2560, 1440
FONT = "Ubuntu Sans"
FONT_PATH = "/usr/share/fonts/truetype/ubuntu/UbuntuSans[wdth,wght].ttf"
QR_URL = "https://huggingface.co/spaces/nourhan44/ferroflux-portal"
QR_TEXT = "huggingface.co/spaces/nourhan44/ferroflux-portal"

BLUE, PURPLE, CYAN = "#16235E", "#7B2FF7", "#00B0D4"
ORANGE, ORANGE_TXT, ORANGE_TINT = "#FF7A2F", "#D9500F", "#FFF1E8"
GREEN, MAGENTA = "#16A34A", "#E8398B"
INK, MUTED, SURFACE, LINE = "#1B2440", "#5C6784", "#F6F8FD", "#DDE3F2"

_fonts = {}


def tw(s, size, weight=400):
    f = _fonts.get((size, weight))
    if f is None:
        f = ImageFont.truetype(FONT_PATH, size)
        f.set_variation_by_axes([100, weight])
        _fonts[(size, weight)] = f
    return f.getlength(s)


out = []


def add(s):
    out.append(s)


def rect(x, y, w, h, r=0, fill="none", stroke=None, sw=1, op=None, extra=""):
    st = f' stroke="{stroke}" stroke-width="{sw}"' if stroke else ""
    o = f' fill-opacity="{op}"' if op is not None else ""
    add(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{r}" fill="{fill}"{o}{st} {extra}/>')


def text(x, y, s, size, weight=400, fill=INK, anchor="start", extra=""):
    add(f'<text x="{x:.1f}" y="{y:.1f}" font-family="{FONT}" font-size="{size}" font-weight="{weight}" '
        f'fill="{fill}" text-anchor="{anchor}" {extra}>{escape(s)}</text>')


def line(x1, y1, x2, y2, stroke, sw=2, extra=""):
    add(f'<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" stroke="{stroke}" stroke-width="{sw}" {extra}/>')


def wrap(s, size, weight, maxw):
    lines, cur = [], ""
    for word in s.split():
        t = (cur + " " + word).strip()
        if tw(t, size, weight) <= maxw or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def rich(x, y, segs, size, maxw, lh):
    """segs: list of (text, weight, fill). Returns y after last line."""
    words = []
    for t, wt, fl in segs:
        for w_ in t.split():
            words.append((w_, wt, fl))
    lines, cur, curw = [], [], 0
    space = tw(" ", size, 400)
    for w_, wt, fl in words:
        ww = tw(w_, size, wt)
        if cur and curw + space + ww > maxw:
            lines.append(cur)
            cur, curw = [], 0
        curw += (space if cur else 0) + ww
        cur.append((w_, wt, fl))
    if cur:
        lines.append(cur)
    for ln in lines:
        spans = ""
        for i, (w_, wt, fl) in enumerate(ln):
            pre = " " if i else ""
            spans += f'<tspan font-weight="{wt}" fill="{fl}">{escape(pre + w_)}</tspan>'
        add(f'<text x="{x:.1f}" y="{y:.1f}" font-family="{FONT}" font-size="{size}" xml:space="preserve">{spans}</text>')
        y += lh
    return y


ICONS = {
    "database": '<ellipse cx="12" cy="6" rx="7" ry="3"/><path d="M5 6v12c0 1.7 3.1 3 7 3s7-1.3 7-3V6"/><path d="M5 12c0 1.7 3.1 3 7 3s7-1.3 7-3"/>',
    "bolt": '<path d="M13 2L4 14h7l-1 8 9-12h-7z"/>',
    "bell": '<path d="M6 8a6 6 0 0112 0c0 7 3 9 3 9H3s3-2 3-9"/><path d="M10.3 21a1.9 1.9 0 003.4 0"/>',
    "layers": '<path d="M12 2l10 5-10 5L2 7z"/><path d="M2 12l10 5 10-5"/><path d="M2 17l10 5 10-5"/>',
    "trend": '<path d="M3 17l6-6 4 4 8-8"/><path d="M15 7h6v6"/>',
    "bars": '<path d="M5 21V10M12 21V3M19 21v-7"/>',
    "stream": '<path d="M3 7h14M13 3l4 4-4 4M21 17H7M11 13l-4 4 4 4"/>',
    "nodes": '<circle cx="5" cy="6" r="2.5"/><circle cx="19" cy="6" r="2.5"/><circle cx="12" cy="18" r="2.5"/><path d="M7.5 6h9M6.5 8l4 8M17.5 8l-4 8"/>',
    "window": '<rect x="2" y="4" width="20" height="16" rx="2"/><path d="M2 9h20"/>',
    "cube": '<path d="M12 2l9 5v10l-9 5-9-5V7z"/><path d="M12 12l9-5M12 12L3 7M12 12v10"/>',
    "spark": '<path d="M12 2l2.2 6.8L21 11l-6.8 2.2L12 20l-2.2-6.8L3 11l6.8-2.2z"/>',
    "link": '<path d="M10 14a5 5 0 007 0l3-3a5 5 0 00-7-7l-1 1M14 10a5 5 0 00-7 0l-3 3a5 5 0 007 7l1-1"/>',
    "shield": '<path d="M12 2l8 3v6c0 5-3.4 9.4-8 11-4.6-1.6-8-6-8-11V5z"/><path d="M9 12l2 2 4-4"/>',
    "upload": '<path d="M12 16V4M7 9l5-5 5 5M4 20h16"/>',
    "userplus": '<circle cx="9" cy="8" r="4"/><path d="M2 21c0-4 3-6 7-6s7 2 7 6M19 8v6M16 11h6"/>',
    "alert": '<path d="M12 3l10 18H2z"/><path d="M12 10v5M12 18v.5"/>',
    "sparkle": '<path d="M12 2v6M12 16v6M2 12h6M16 12h6M5 5l4 4M15 15l4 4M19 5l-4 4M9 15l-4 4"/>',
}


def _load_brands():
    out_ = {}
    for nm in ("apachekafka", "apachespark", "apacheairflow", "postgresql", "fastapi", "n8n", "streamlit", "docker", "huggingface", "neon"):
        with open(os.path.join(HERE, "assets", "icons", nm + ".svg"), encoding="utf-8") as f:
            out_[nm] = re.findall(r' d="([^"]+)"', f.read())
    return out_


BRANDS = _load_brands()


def brand(name, x, y, size, fill=INK, op=1):
    sc = size / 24
    ps = "".join(f'<path d="{d}"/>' for d in BRANDS[name])
    add(f'<g transform="translate({x:.1f} {y:.1f}) scale({sc:.3f})" fill="{fill}" fill-opacity="{op}">{ps}</g>')


def brand_icon(name, x, y, size, color):
    add(f'<rect x="{x:.1f}" y="{y:.1f}" width="{size}" height="{size}" rx="{size * 0.26:.1f}" fill="{color}" fill-opacity="0.13"/>')
    inner_ = size * 0.58
    brand(name, x + (size - inner_) / 2, y + (size - inner_) / 2, inner_, color)


def icon(name, x, y, size, color, stroke_scale=2.0):
    add(f'<rect x="{x:.1f}" y="{y:.1f}" width="{size}" height="{size}" rx="{size * 0.26:.1f}" fill="{color}" fill-opacity="0.13"/>')
    pad = size * 0.2
    sc = (size - 2 * pad) / 24
    add(f'<g transform="translate({x + pad:.1f} {y + pad:.1f}) scale({sc:.3f})" fill="none" stroke="{color}" '
        f'stroke-width="{stroke_scale}" stroke-linecap="round" stroke-linejoin="round">{ICONS[name]}</g>')


def shadow(x, y, w, h, r):
    for d, o in ((7, 0.025), (5, 0.03), (3, 0.04), (1.5, 0.05)):
        rect(x - d / 3, y + d, w + 2 * d / 3, h, r + 2, "#16235E", op=o)


def card(x, y, w, h, n, title, color, accent=None):
    shadow(x, y, w, h, 22)
    rect(x, y, w, h, 22, SURFACE, LINE, 1.5)
    if accent:
        rect(x, y + 18, 6, h - 36, 3, accent)
    add(f'<circle cx="{x + 40}" cy="{y + 36}" r="20" fill="{color}"/>')
    text(x + 40, y + 44, str(n), 22, 700, "#FFFFFF", "middle")
    text(x + 72, y + 46, title, 28, 700, INK)
    line(x + 72, y + 56, x + 72 + tw(title, 28, 700), y + 56, color, 3)
    return y + 82


def chip_w(label, size=20, weight=500, pad=20, logo=None, hot=False):
    return tw(label, size, 700 if hot else weight) + 2 * pad + (30 if logo else 0)


def chip(x, y, label, size=20, weight=500, hot=False, h=46, pad=20, logo=None):
    wd = chip_w(label, size, weight, pad, logo, hot)
    if hot:
        rect(x, y, wd, h, 14, ORANGE_TINT, ORANGE, 2.5)
    else:
        rect(x, y, wd, h, 14, "#FFFFFF", LINE, 1.5)
    tx = x + pad
    if logo:
        brand(logo, tx, y + (h - 22) / 2, 22, INK, 0.88)
        tx += 30
    text(tx, y + h / 2 + size * 0.34, label, size, 700 if hot else weight, INK)
    return wd


# ------------------------------------------------------------------ defs / page
def logo(name):
    im = Image.open(os.path.join(HERE, "assets", name))
    if im.mode != "RGBA":
        im = im.convert("RGBA")
        r_, g_, b_, _a = im.split()
        mn = ImageChops.darker(ImageChops.darker(r_, g_), b_)
        im.putalpha(mn.point(lambda v: 255 if v < 205 else max(0, int((250 - v) * 255 / 45))))
    bb = im.getchannel("A").getbbox()
    im = im.crop(bb)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return base64.b64encode(buf.getvalue()).decode(), im.size[0] / im.size[1]


add(f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="{W}" height="{H}" viewBox="0 0 {W} {H}">')
add(f'''<defs>
<linearGradient id="gt" gradientUnits="userSpaceOnUse" x1="1160" y1="0" x2="1660" y2="0">
<stop offset="0" stop-color="{BLUE}"/><stop offset="0.5" stop-color="{PURPLE}"/><stop offset="1" stop-color="{MAGENTA}"/></linearGradient>
<pattern id="dots" width="28" height="28" patternUnits="userSpaceOnUse"><circle cx="2" cy="2" r="1.3" fill="#C5CFEA"/></pattern>
<radialGradient id="glow" cx="0.5" cy="0" r="0.75"><stop offset="0" stop-color="#7B2FF7" stop-opacity="0.13"/><stop offset="1" stop-color="#7B2FF7" stop-opacity="0"/></radialGradient>
<radialGradient id="fglow" cx="0.85" cy="0.1" r="0.7"><stop offset="0" stop-color="#FFFFFF" stop-opacity="0.16"/><stop offset="1" stop-color="#FFFFFF" stop-opacity="0"/></radialGradient>
<linearGradient id="gf" x1="0" y1="0" x2="1" y2="0"><stop offset="0" stop-color="{BLUE}"/><stop offset="1" stop-color="{PURPLE}"/></linearGradient>
<linearGradient id="gl" x1="0" y1="0" x2="1" y2="0"><stop offset="0" stop-color="{BLUE}"/><stop offset="0.55" stop-color="{PURPLE}"/><stop offset="1" stop-color="{CYAN}"/></linearGradient>
</defs>''')
rect(0, 0, W, H, 0, "#FFFFFF")
rect(0, 0, W, H, 0, "url(#dots)", op=0.5)
rect(0, 0, W, 560, 0, "url(#glow)")

# ------------------------------------------------------------------ header
HDR = 198
LOGO_H = 112
ly = (HDR - 6 - LOGO_H) / 2
x_ = 48
for nm, hh in (("LogoMinistry.png", 100), ("LogoMilitary.png", LOGO_H)):
    d_, ar = logo(nm)
    add(f'<image x="{x_}" y="{(HDR - 6 - hh) / 2:.1f}" width="{hh * ar:.1f}" height="{hh}" xlink:href="data:image/png;base64,{d_}"/>')
    x_ += hh * ar + 32
d_, ar = logo("LogoDigilians.jpg")
add(f'<image x="{W - 48 - LOGO_H * ar:.1f}" y="{ly:.1f}" width="{LOGO_H * ar:.1f}" height="{LOGO_H}" xlink:href="data:image/png;base64,{d_}"/>')
MK_H = 90
mark_b64 = base64.b64encode(open(os.path.join(HERE, "assets", "logo_mark.png"), "rb").read()).decode()
mk_w = MK_H * 276 / 282
title_w = tw("FerroFlux", 92, 700)
blk_w = mk_w + 18 + title_w
bx0 = W / 2 - blk_w / 2
add(f'<image x="{bx0:.1f}" y="10" width="{mk_w:.1f}" height="{MK_H}" xlink:href="data:image/png;base64,{mark_b64}"/>')
text(bx0 + mk_w + 18, 92, "FerroFlux", 92, 700, "url(#gt)", extra="")
text(W / 2, 136, "Intelligent Steel Supply-Chain Analytics Platform", 34, 700, INK, "middle")
text(W / 2, 176, "Digilians · Nano Degree (4 Months) · Track: AI & Data Science — Professional Big Data Engineer",
     27, 600, MUTED, "middle")
for lab, side in (("Open source · free-tier cloud", "L"), ("Think Steel. Think Smart.", "R")):
    pw_ = tw(lab, 22, 700) + 44
    px_ = (bx0 - 28 - pw_) if side == "L" else (bx0 + blk_w + 28)
    rect(px_, 50, pw_, 42, 21, "#FFFFFF", PURPLE, 2.5)
    text(px_ + pw_ / 2, 78, lab, 22, 700, PURPLE, "middle")
rect(0, HDR - 6, W, 6, 0, "url(#gl)")

# ------------------------------------------------------------------ geometry
M = 48
TOP = HDR + 18
FOOT_H = 150
FOOT_Y = H - M - FOOT_H
BOT = FOOT_Y - 20
GAP = 20
LW = 592
CX = M + LW + 28
CW = W - 2 * M - 2 * LW - 56
RX = W - M - LW
inner = LW - 48

# ------------------------------------------------------------------ LEFT column
lh_ = [236, 206, 262, 242]
y = TOP
# 1 identity
cy = card(M, y, LW, lh_[0], 1, "Project Identity", BLUE)
rows = [("Project", "FerroFlux Platform"), ("Supervisor", "Dr. Rania ElGohary"),
        ("Track Head", "Dr. Mohamed El Badawi"), ("Instructor", "Dr. Amera Youssef")]
yy = cy + 22
for k, v in rows:
    text(M + 28, yy, k, 18, 500, MUTED)
    text(M + 190, yy, v, 22, 700, INK)
    yy += 34
y += lh_[0] + GAP
# 2 team
cy = card(M, y, LW, lh_[1], 2, "Team", PURPLE)
names = ["Nourhan Saber Shalaby", "Maram Mostafa", "Hager Elkahlawy", "Zeinab Elgabalawy", "Haidy Nageh"]
for i, nm in enumerate(names):
    col, row = i % 2, i // 2
    px = M + 28 + col * 272
    py = cy + 24 + row * 36
    add(f'<circle cx="{px + 6}" cy="{py - 7}" r="5" fill="{PURPLE}"/>')
    text(px + 22, py, nm, 21, 600, INK)
y += lh_[1] + GAP
# 3 problem & audience
cy = card(M, y, LW, lh_[2], 3, "Problem & Audience", CYAN)
tw3 = (inner - 24) / 3
for i, (num, lab) in enumerate((("86,044", "rows of data"), ("3", "factories"), ("12", "suppliers"))):
    tx = M + 24 + i * (tw3 + 12)
    rect(tx, cy, tw3, 70, 14, "#FFFFFF", LINE, 1.5)
    text(tx + tw3 / 2, cy + 34, num, 34, 700, BLUE, "middle")
    text(tx + tw3 / 2, cy + 58, lab, 17, 500, MUTED, "middle")
text(M + 28, cy + 92, "Simulated EZZ Steel Group dataset", 16, 500, MUTED)
yy = cy + 124
for ln in wrap("Fragmented spreadsheets and manual price tracking", 21, 600, inner):
    text(M + 28, yy, ln, 21, 600, INK)
    yy += 28
yy += 4
for ln in wrap("Users: factory managers, admins, decision makers", 21, 400, inner):
    text(M + 28, yy, ln, 21, 400, MUTED)
    yy += 28
y += lh_[2] + GAP
# 4 solution
cy = card(M, y, LW, lh_[3], 4, "Solution & Core Idea", BLUE, accent=ORANGE)
rich(M + 28, cy + 26, [("One platform that streams, cleans and predicts — and", 600, INK),
                       ("automatically alerts each factory", 700, ORANGE_TXT),
                       ("to price spikes and low efficiency.", 600, INK)], 24, inner - 6, 34)

# ------------------------------------------------------------------ RIGHT column
rh = [196, 244, 276, 226]
y = TOP
# 5 pipeline
cy = card(RX, y, LW, rh[0], 5, "Pipeline", PURPLE)
steps = [("brand:apachekafka", ("Ingest", "Kafka"), PURPLE, False), ("brand:apachespark", ("Clean", "Silver"), PURPLE, False),
         ("bell", ("Detect &", "Alert"), ORANGE, True), ("brand:postgresql", ("Aggregate", "Gold x6"), PURPLE, False),
         ("trend", ("Predict", "3 models"), PURPLE, False)]
pw = inner / 5
for i, (ic, lab, col, hot) in enumerate(steps):
    cxs = RX + 24 + pw * i + pw / 2
    (brand_icon(ic[6:], cxs - 24, cy + 2, 48, col) if ic.startswith('brand:') else icon(ic, cxs - 24, cy + 2, 48, col))
    text(cxs, cy + 76, lab[0], 18, 700, ORANGE_TXT if hot else INK, "middle")
    text(cxs, cy + 97, lab[1], 18 if hot else 16, 700 if hot else 500, ORANGE_TXT if hot else MUTED, "middle")
    if i < 4:
        ax = cxs + pw / 2
        add(f'<path d="M{ax - 9:.1f} {cy + 26} h14 m-5 -5 l5 5 l-5 5" fill="none" stroke="{MUTED}" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round"/>')
y += rh[0] + GAP
# 6 key features
cy = card(RX, y, LW, rh[1], 6, "Key Features", CYAN)
feats = [("shield", "Tenant isolation", "company_id on every read", BLUE, False),
         ("upload", "Excel to auto-ETL", "upload triggers Airflow", CYAN, False),
         ("userplus", "Self-service signup", "register, reset by email", PURPLE, False),
         ("alert", "Spike alerts", "email, inbox, n8n", BLUE, True),
         ("trend", "3 ML forecasts", "price, demand, risk", PURPLE, False),
         ("stream", "Live streaming", "Kafka to Spark", CYAN, False)]
cwf = inner / 2
for i, (ic, lab, sub, col, hot) in enumerate(feats):
    fx = RX + 24 + (i % 2) * cwf
    fy = cy + (i // 2) * 52
    icon(ic, fx, fy, 42, col)
    text(fx + 54, fy + 19, lab, 18, 700, ORANGE_TXT if hot else INK)
    text(fx + 54, fy + 38, sub, 15, 500, MUTED)
y += rh[1] + GAP
# 7 competitive matrix
cy = card(RX, y, LW, rh[2], 7, "Why FerroFlux", BLUE)
lc = 190
cc = (inner - lc) / 4
tx0 = RX + 24
rect(tx0 + lc, cy - 4, cc, 36 + 6 * 25 + 4, 10, BLUE, op=0.10)
for j, hd in enumerate((["FerroFlux"], ["SAP", "IBP"], ["Tableau", "+ Excel"], ["Generic", "BI"])):
    for k, ln_ in enumerate(hd):
        yy_ = cy + 14 + k * 16 if len(hd) == 2 else cy + 22
        text(tx0 + lc + cc * j + cc / 2, yy_, ln_, 15, 700, BLUE if j == 0 else MUTED, "middle")


def sym(kind, x, y_):
    if kind == "y":
        add(f'<path d="M{x - 7} {y_} l5 5 l10 -11" fill="none" stroke="{GREEN}" stroke-width="3.2" stroke-linecap="round" stroke-linejoin="round"/>')
    elif kind == "n":
        add(f'<path d="M{x - 5} {y_ - 5} l10 10 M{x + 5} {y_ - 5} l-10 10" fill="none" stroke="#9AA3BB" stroke-width="2.6" stroke-linecap="round"/>')
    else:
        add(f'<circle cx="{x}" cy="{y_}" r="6.5" fill="none" stroke="{MUTED}" stroke-width="2"/>')
        add(f'<path d="M{x} {y_ - 6.5} a6.5 6.5 0 0 1 0 13z" fill="{MUTED}"/>')


mrows = [("Real-time streaming", "ynnp"), ("Built-in ML forecasts", "yynp"), ("Automated alerts", "yynp"),
         ("Steel-industry focus", "ynnn"), ("Self-service signup", "ynnp")]
ry0 = cy + 36
for i, (lab, vals) in enumerate(mrows):
    ry = ry0 + i * 25
    if i:
        line(tx0, ry - 5, tx0 + inner, ry - 5, LINE, 1)
    text(tx0 + 4, ry + 14, lab, 16, 600, INK)
    for j, v in enumerate(vals):
        sym(v, tx0 + lc + cc * j + cc / 2, ry + 9)
ry = ry0 + 5 * 25
line(tx0, ry - 5, tx0 + inner, ry - 5, LINE, 1)
text(tx0 + 4, ry + 14, "Deployment cost", 16, 600, INK)
for j, v in enumerate(("Free tier", "Very high", "Medium", "High")):
    text(tx0 + lc + cc * j + cc / 2, ry + 14, v, 15, 700 if j == 0 else 500, GREEN if j == 0 else MUTED, "middle")
y += rh[2] + GAP
# 8 KPIs
cy = card(RX, y, LW, rh[3], 8, "Impact & KPIs", GREEN)
kp = [("R² 0.9578", "Price forecast R²"), ("MAPE 1.61%", "Price forecast error"),
      ("94.78%", "Supplier-risk accuracy"), ("R² 0.80", "Demand forecast R²")]
kw = (inner - 12) / 2
for i, (num, lab) in enumerate(kp):
    kx = RX + 24 + (i % 2) * (kw + 12)
    ky = cy + (i // 2) * 72
    rect(kx, ky, kw, 66, 14, "#FFFFFF", LINE, 1.5)
    text(kx + 18, ky + 32, num, 32, 700, GREEN)
    text(kx + 18, ky + 54, lab, 16, 500, MUTED)

# ------------------------------------------------------------------ CENTRE: phones + architecture
SW, SH = 264, 571
BEZ = 10


def phone_img(name, y0=0):
    im = Image.open(os.path.join(HERE, "assets", "shots", "mobile_live", name)).convert("RGB")
    im = im.crop((0, y0, 1170, y0 + 2532)).resize((SW * 3, SH * 3), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=93, subsampling=0)
    return base64.b64encode(buf.getvalue()).decode()


PANEL_H = 664
add('<defs><linearGradient id="gp" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#E7ECFB"/><stop offset="1" stop-color="#F6F8FD"/></linearGradient></defs>')
shadow(CX, TOP, CW, PANEL_H, 24)
rect(CX, TOP, CW, PANEL_H, 24, "url(#gp)", LINE, 1.5)
phones = [("a_overview.png", 0, "1", "Tenant overview"), ("a_spikes.png", 0, "2", "Spike & demand analytics"),
          ("a_suppliers.png", 0, "3", "Supplier risk scorecard"), ("a_upload.png", 0, "4", "Excel upload to ETL")]
pw_o = SW + 2 * BEZ
gapp = (CW - 4 * pw_o) / 5
for i, (nm, y0, num, cap) in enumerate(phones):
    px = CX + gapp + i * (pw_o + gapp)
    py = TOP + 16
    ph_h = SH + 2 * BEZ
    shadow(px, py, pw_o, ph_h, 42)
    rect(px, py, pw_o, ph_h, 42, "#0E1424", "#2A3350", 2.5)
    add(f'<clipPath id="ps{i}"><rect x="{px + BEZ}" y="{py + BEZ}" width="{SW}" height="{SH}" rx="32"/></clipPath>')
    add(f'<image x="{px + BEZ}" y="{py + BEZ}" width="{SW}" height="{SH}" clip-path="url(#ps{i})" preserveAspectRatio="none" '
        f'xlink:href="data:image/jpeg;base64,{phone_img(nm, y0)}"/>')
    rect(px + pw_o / 2 - 34, py + BEZ + 9, 68, 19, 9.5, "#0E1424")
    rect(px - 3, py + 120, 3, 46, 1.5, "#2A3350")
    rect(px + pw_o, py + 150, 3, 70, 1.5, "#2A3350")
    cy_cap = py + ph_h + 14
    capw = tw(cap, 18, 700) + 64
    rect(px + pw_o / 2 - capw / 2, cy_cap, capw, 34, 17, "#FFFFFF", LINE, 1.5)
    add(f'<circle cx="{px + pw_o / 2 - capw / 2 + 19}" cy="{cy_cap + 17}" r="11" fill="{BLUE}"/>')
    text(px + pw_o / 2 - capw / 2 + 19, cy_cap + 23, num, 15, 700, "#FFFFFF", "middle")
    text(px + pw_o / 2 - capw / 2 + 38, cy_cap + 23, cap, 18, 700, INK)

TIER_H, CHEV = 60, 12
strip_y = TOP + PANEL_H + 14
text(CX + 4, strip_y + 20, "SYSTEM ARCHITECTURE", 19, 700, MUTED, extra='letter-spacing="2"')
aw = tw("Alert", 26, 700)
text(CX + CW - 4, strip_y + 22, "Alert", 26, 700, ORANGE_TXT, "end")
text(CX + CW - 4 - aw - tw(" ", 26, 700), strip_y + 22, "Stream · Predict ·", 26, 700, INK, "end")
t_top = strip_y + 34
tiers = [t_top + i * (TIER_H + CHEV) for i in range(4)]
LBL_W = 214
tier_x, tier_w = CX, CW
lane_x = tier_x + tier_w - 28 - 70
tcolors = [CYAN, BLUE, PURPLE, INK]
tlabels = [("1", "SOURCES"), ("2", "STREAM & ETL"), ("3", "STORE & ML"), ("4", "SERVE & ALERT")]
for i, ty in enumerate(tiers):
    rect(tier_x, ty, tier_w, TIER_H, 18, SURFACE, LINE, 1.5)
    add(f'<path d="M{tier_x + 18} {ty} h{LBL_W - 18} v{TIER_H} h-{LBL_W - 18} a18 18 0 0 1 -18 -18 v-{TIER_H - 36} a18 18 0 0 1 18 -18 z" fill="{tcolors[i]}"/>')
    text(tier_x + 24, ty + 42, tlabels[i][0], 30, 700, "#FFFFFF", extra='opacity="0.55"')
    text(tier_x + 56, ty + 38, tlabels[i][1], 20, 700, "#FFFFFF")
    if i < 3:
        mx = tier_x + tier_w / 2
        add(f'<path d="M{mx - 10:.1f} {ty + TIER_H + 1} l10 9 l10 -9" fill="none" stroke="{MUTED}" stroke-width="3.2" stroke-linecap="round" stroke-linejoin="round"/>')

CH = 38


def chips_row(x, y, items, limit, pad=18):
    for it in items:
        lb, lg = (it, None) if isinstance(it, str) else it
        wd = chip_w(lb, 19, 500, pad, lg)
        assert x + wd <= limit, (lb, x + wd, limit)
        chip(x, y, lb, size=19, h=CH, pad=pad, logo=lg)
        x += wd + 12


cxs0 = tier_x + LBL_W + 24
lim = lane_x - 30
off = (TIER_H - CH) / 2
right_lim = tier_x + tier_w - 24
chips_row(cxs0, tiers[0] + off, ["Factory telemetry", "Market price feed", "Excel uploads"], right_lim)
hot2, hot4 = "Anomaly detection", "Email + portal alerts"
h2w, h4w = chip_w(hot2, 19, 500, 18, None, True), chip_w(hot4, 19, 500, 18, None, True)
chips_row(cxs0, tiers[1] + off, [("Apache Kafka", "apachekafka"), ("Spark: Bronze › Silver › Gold", "apachespark"), ("Airflow", "apacheairflow")], right_lim - h2w - 14)
chips_row(cxs0, tiers[2] + off, [("PostgreSQL", "postgresql"), ("Neon", "neon"), "Price forecast", "Demand forecast", "Supplier risk"], lim, pad=15)
chips_row(cxs0, tiers[3] + off, [("FastAPI", "fastapi"), ("Streamlit", "streamlit"), ("n8n", "n8n"), ("Docker", "docker"), ("Hugging Face", "huggingface")], right_lim - h4w - 14, pad=15)
for tidx, lab in ((1, hot2), (3, hot4)):
    wd = chip_w(lab, 19, 500, 18, None, True)
    chip(right_lim - wd, tiers[tidx] + off, lab, size=19, h=CH, pad=18, hot=True)
y1 = tiers[1] + off + CH
y2 = tiers[3] + off
add(f'<line x1="{lane_x:.1f}" y1="{y1:.1f}" x2="{lane_x:.1f}" y2="{y2 - 8:.1f}" stroke="{ORANGE}" stroke-width="4" stroke-dasharray="9 7" stroke-linecap="round"/>')
add(f'<path d="M{lane_x - 9:.1f} {y2 - 16:.1f} l9 11 l9 -11" fill="none" stroke="{ORANGE}" stroke-width="4" stroke-linecap="round" stroke-linejoin="round"/>')
lm = tiers[2] + TIER_H / 2
rect(lane_x - 50, lm - 13, 100, 26, 8, SURFACE)
text(lane_x, lm + 6, "ALERT PATH", 15, 700, ORANGE_TXT, "middle")

# ------------------------------------------------------------------ footer
shadow(M, FOOT_Y, W - 2 * M, FOOT_H, 28)
rect(M, FOOT_Y, W - 2 * M, FOOT_H, 28, "url(#gf)")
fy = FOOT_Y + FOOT_H / 2
rect(M, FOOT_Y, W - 2 * M, FOOT_H, 28, "url(#fglow)")
lf_b64 = base64.b64encode(open(os.path.join(HERE, "assets", "logo_full.png"), "rb").read()).decode()
rect(M + 24, fy - 62, 142, 124, 20, "#FFFFFF")
add(f'<image x="{M + 33}" y="{fy - 54}" width="{124}" height="{124 * 367 / 424:.1f}" xlink:href="data:image/png;base64,{lf_b64}"/>')
add(f'<g transform="translate({M + 188} {fy - 22}) scale(1.83)" fill="none" stroke="#FFFFFF" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">{ICONS["link"]}</g>')
text(M + 252, fy - 2, "PROJECT LINK", 30, 700, "#FFFFFF")
text(M + 252, fy + 34, "& QR CODE", 30, 700, "#FFFFFF")

qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_H, border=0, box_size=1)
qr.add_data(QR_URL)
qr.make(fit=True)
mat = qr.get_matrix()
n = len(mat)
ms = 3
qsz = n * ms
plate = qsz + 14
px0 = 880
py0 = FOOT_Y + (FOOT_H - plate) / 2
rect(px0, py0, plate, plate, 12, "#FFFFFF")
qx, qy = px0 + 7, py0 + 7
for r_, row in enumerate(mat):
    c_ = 0
    while c_ < n:
        if row[c_]:
            s = c_
            while c_ < n and row[c_]:
                c_ += 1
            rect(qx + s * ms, qy + r_ * ms, (c_ - s) * ms, ms, 0, BLUE, extra='shape-rendering="crispEdges"')
        else:
            c_ += 1
text(px0 + plate + 32, fy - 4, "Live portal · scan or open", 22, 500, "#D6DCF5")
text(px0 + plate + 32, fy + 36, QR_TEXT, 31, 700, "#FFFFFF")

right = W - M - 40
rect(right - 100, fy - 50, 100, 100, 50, "#FFFFFF")
d_, ar = logo("LogoMilitary.png")
add(f'<image x="{right - 92}" y="{fy - 42}" width="84" height="{84 / ar:.1f}" xlink:href="data:image/png;base64,{d_}"/>')
text(right - 124, fy - 2, "Digilians · Graduation Project", 27, 700, "#FFFFFF", "end")
text(right - 124, fy + 34, "Digital Pioneers Initiative · 2026", 22, 500, "#D6DCF5", "end")

add("</svg>")
svg = "\n".join(out)
with open(os.path.join(OUT, "FerroFlux_Infographic.svg"), "w", encoding="utf-8") as f:
    f.write(svg)
b_ = svg.encode("utf-8")
png = os.path.join(OUT, "FerroFlux_Infographic_4K.png")
cairosvg.svg2png(bytestring=b_, write_to=os.path.join(OUT, "FerroFlux_Infographic.png"), output_width=W, output_height=H)
cairosvg.svg2png(bytestring=b_, write_to=png, output_width=3840, output_height=2160)
cairosvg.svg2png(bytestring=b_, write_to=os.path.join(OUT, "FerroFlux_Infographic_5K_print.png"), output_width=5120, output_height=2880)
cairosvg.svg2pdf(bytestring=b_, write_to=os.path.join(OUT, "FerroFlux_Infographic.pdf"), output_width=W, output_height=H)

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
sl = prs.slides.add_slide(prs.slide_layouts[6])
sl.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
prs.save(os.path.join(OUT, "FerroFlux_Infographic.pptx"))
print("built", n, "modules QR")
